import os
import shutil
import subprocess
import platform
from openpyxl import load_workbook

template_file = "TEMPLATE.xlsx" # Excel template for data logging
base_name = "prefix_here_" 
ext = ".xlsx"

# Find next available filename

def save_report(save_path, ramp_up_current_Max, ramp_up_current, ramp_up_passFail,
                positive_5V_min, positive_5V_max, positive_5V_voltage, positive_5V_passFail,
                negative_5V_min, negative_5V_max, negative_5V_voltage, negative_5V_passFail,
                CLKIN_LOW_VOLTAGE_MAX, CLKIN_LOW_OUTPUT_voltage, CLKIN_LOW_OUTPUT_passFail,
                CLKIN_HIGH_CURRENT_MAX, CLKIN_HIGH_CURRENT, CLKIN_HIGH_CURRENT_passFail,
                CLKIN_HIGH_OUTPUT_voltage, CLKIN_HIGH_OUTPUT_passFail, positive_5V_CLKIN_max, positive_5V_CLKIN_min,
                positive_5V_DATAOUT_max, positive_5V_DATAOUT_min, positive_5V_DATAOUT_voltage, positive_5V_DATAOUT_passFail,
                DATAIN_LOW_OUTPUT_voltage, DATAIN_LOW_OUTPUT_passFail, DATAIN_HIGH_CURRENT, DATAIN_HIGH_CURRENT_passFail,
                DATAIN_HIGH_OUTPUT_voltage, DATAIN_HIGH_OUTPUT_passFail, DATAIN_HIGH_CURRENT_MAX,
                DATAIN_HIGH_VOLTAGE_MAX, positive_5V_DATAIN_max, positive_5V_DATAIN_min,
                AUXIN_max, AUXIN_min, AUXIN_voltage, AUXIN_passFail):

    counter = 1
    while True:
        DOC_SN = f"{base_name}{counter:05d}{ext}"
        new_filename = f"{save_path}{DOC_SN}"
        if not os.path.exists(new_filename):
            break
        counter += 1

    shutil.copy(template_file, new_filename) # Copy template to new file
    wb = load_workbook(new_filename) # Open the copied file and modify it
    ws = wb.active  # or wb["SheetName"] if needed

    ########################### FILE NAME ###########################
    ws["H12"] = DOC_SN # DOC S/N

    ########################### RAMP UP TEST ###########################
    ws["D27"] = f"<{ramp_up_current_Max}A" # expected value
    ws["G27"] = f"{ramp_up_current}A"
    ws["I27"] = ramp_up_passFail

    ########################### +5V TEST ###########################
    ws["D29"] = f"{positive_5V_min}V to {positive_5V_max}V" # expected value
    ws["G29"] = f"{positive_5V_voltage:.3f}V" # measured value
    ws["I29"] = positive_5V_passFail # pass fail

    ########################### -5V TEST ###########################
    ws["D31"] = f"{negative_5V_min}V to {negative_5V_max}V" # expected value
    ws["G31"] = f"{negative_5V_voltage:.3f}V" # measured value
    ws["I31"] = negative_5V_passFail # pass fail

    ########################### CLKIN LOW Vout (R9) TEST ###########################
    ws["D33"] = f"<{CLKIN_LOW_VOLTAGE_MAX}V" # expected value
    ws["G33"] = f"{CLKIN_LOW_OUTPUT_voltage:.3f}V" # measured value
    ws["I33"] = CLKIN_LOW_OUTPUT_passFail # pass fail

    ########################### CLKIN HIGH CURRENT TEST ###########################
    ws["D35"] = f"<{CLKIN_HIGH_CURRENT_MAX}A" # expected value
    ws["G35"] = f"{CLKIN_HIGH_CURRENT:.3f}A" # measured value
    ws["I35"] = CLKIN_HIGH_CURRENT_passFail # pass fail

    # ########################### CLKIN HIGH Vout (R9) TEST ###########################
    ws["D37"] = f"{positive_5V_CLKIN_min}V to {positive_5V_CLKIN_max}V" # expected value
    ws["G37"] = f"{CLKIN_HIGH_OUTPUT_voltage:.3f}V" # measured value
    ws["I37"] = CLKIN_HIGH_OUTPUT_passFail # pass fail

    # ########################### DATA OUT + TEST ###########################
    ws["D39"] = f"{positive_5V_DATAOUT_min}V to {positive_5V_DATAOUT_max}V" # expected value
    ws["G39"] = f"{positive_5V_DATAOUT_voltage:.3f}V" # measured value
    ws["I39"] = positive_5V_DATAOUT_passFail # pass fail

    ########################### DATAIN HIGH CURRENT TEST ###########################
    ws["D41"] = f"<{DATAIN_HIGH_CURRENT_MAX}A" # expected value
    ws["G41"] = f"{DATAIN_HIGH_CURRENT:.3f}A" # measured value
    ws["I41"] = DATAIN_HIGH_CURRENT_passFail # pass fail

    ########################### DATAIN HIGH Vout (Pin9) TEST ###########################
    ws["D43"] = f"<{DATAIN_HIGH_VOLTAGE_MAX}V"  # expected value
    ws["G43"] = f"{DATAIN_HIGH_OUTPUT_voltage:.3f}V" # measured value
    ws["I43"] = DATAIN_HIGH_OUTPUT_passFail # pass fail

    # ########################### DATAIN LOW Vout (Pin9) TEST ###########################
    ws["D45"] = f"{positive_5V_DATAIN_min}V to {positive_5V_DATAIN_max}V"  # expected value
    ws["G45"] = f"{DATAIN_LOW_OUTPUT_voltage:.3f}V" # measured value
    ws["I45"] = DATAIN_LOW_OUTPUT_passFail # pass fail

    # ############################## AUXIN (Pin8) TEST #################################
    ws["D47"] = f"{AUXIN_min}V to {AUXIN_max}V"  # expected value
    ws["G47"] = f"{AUXIN_voltage:.3f}V" # measured value
    ws["I47"] = AUXIN_passFail # pass fail


    # Save changes
    wb.save(new_filename)
    print(f"Created file: {new_filename}")

    ####################################################################
    # OPEN WITH LIBREOFFICE
    ####################################################################


    try:
        system = platform.system()

        if system == "Windows":
            subprocess.Popen([
                r"C:\Program Files\LibreOffice\program\soffice.exe",
                new_filename
            ])

        elif system == "Linux":
            subprocess.Popen(["libreoffice", new_filename])

        elif system == "Darwin":  # macOS
            subprocess.Popen(["open", "-a", "LibreOffice", new_filename])

    except Exception as e:
        print(f"Could not open LibreOffice: {e}")



